import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok\Seosm")
OUT = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok\inspection\before_preview_rows.json")


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


result = []
for index, path in enumerate(sorted(ROOT.glob("*.xlsx")), start=1):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    rows = []
    preview_end = min(max_row(path), 20)
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=6, max_row=preview_end, min_col=1, max_col=11, values_only=True),
        start=6,
    ):
        rows.append([row_number, str(values[0]), str(values[2]), str(values[5]), values[10]])
    result.append({"index": index, "file": path.name, "rows": rows})

OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)

