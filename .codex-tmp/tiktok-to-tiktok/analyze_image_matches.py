import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok")


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(value) for value in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


def image_ids(values):
    found = set()
    for value in values:
        found.update(re.findall(r"/([0-9a-f]{32})(?:~|\?|\"|$)", str(value or "").lower()))
    return found


def load(store):
    products = {}
    for path in sorted((ROOT / store).glob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["Template"]
        for row in sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=29, values_only=True):
            pid = str(row[0])
            if pid in products:
                continue
            products[pid] = {
                "name": str(row[2] or ""),
                "file": path.name,
                "ids": image_ids([row[6], *row[17:27]]),
                "main_ids": image_ids(row[17:27]),
                "urls": row[27:29],
            }
    return products


vilion = load("Vilion")
seosm = load("Seosm")
for tid, target in seosm.items():
    print(f"\nTARGET {tid} | images={len(target['ids'])} | {target['name']}")
    scored = []
    for sid, source in vilion.items():
        overlap = target["ids"] & source["ids"]
        main_overlap = target["main_ids"] & source["main_ids"]
        scored.append((len(overlap), len(main_overlap), sid, source, sorted(overlap)))
    scored.sort(reverse=True)
    for total, main, sid, source, overlap in scored[:5]:
        print(f"  overlap={total} main={main} | {sid} | {source['name']} | {overlap[:5]}")
