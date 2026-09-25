import csv
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok")
INSPECTION = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok\inspection")

TARGET_FAMILIES = {
    "1734273764896900728": "JAS-P",
    "1734273792674989688": "3IN1-P",
    "1734273934335313528": "3IN1-P",
    "1734273961869411960": "JC-P",
    "1734273999090583160": "JAS-P",
    "1734274000206268024": "RMP-P",
    "1734274041964365432": "JC-P",
    "1734274054262457976": "RMP-P",
    "1734274070804727416": "3IN1-P",
    "1734274153637054072": "3IN1-P",
    "1734274163153602168": "3IN1-P",
    "1736586260006274680": "KMJ",
    "1734273477733418616": "JAS-W",
    "1734274154383574648": "JASKO",
    "1734274159432205944": "JAS-A",
}

FAMILY_PREFIXES = [
    ("3in1-p-", "3IN1-P"),
    ("jas-p-", "JAS-P"),
    ("rmp-p-", "RMP-P"),
    ("rc-p-", "RC-P"),
    ("cln-p-", "CLN-P"),
    ("jas-w-", "JAS-W"),
    ("3in1-a-", "3IN1-A"),
    ("kmj-", "KMJ"),
    ("dsi-", "DSI"),
]

TARGET_COLOR_ALIASES = {
    "3IN1-P": {
        "biru muda": "biru langit",
        "merah marun": "marun muda",
        "merah": "merah cabe",
    },
    "RMP-P": {
        "abu muda": "abu",
        "marun": "marun muda",
    },
    "JAS-W": {
        "teracota": "terracota",
        "pink": "baby pink",
        "cream": "krem",
        "marun": "marun muda",
        "biru muda": "biru langit",
        "biru tua": "biru",
    },
}


def real_max_row(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(value) for value in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


def clean_part(value: str) -> str:
    return " ".join(re.sub(r"\s+", " ", str(value or "").strip().lower()).split())


def normalized_variation(value: str, family: str, target: bool = False) -> tuple[str, ...]:
    parts = [clean_part(part) for part in str(value or "").split(",")]
    if target and parts:
        parts[0] = TARGET_COLOR_ALIASES.get(family, {}).get(parts[0], parts[0])
    if len(parts) > 1 and parts[-1] == "3xl":
        parts[-1] = "xxxl"
    return tuple(parts)


def sku_family(sku: str) -> str | None:
    value = str(sku or "").strip().lower()
    for prefix, family in FAMILY_PREFIXES:
        if value.startswith(prefix):
            return family
    return None


def sku_tail(sku: str, family: str) -> tuple[str, str] | None:
    value = str(sku).strip().lower()
    prefix = next((prefix for prefix, name in FAMILY_PREFIXES if name == family and value.startswith(prefix)), None)
    if not prefix:
        return None
    tail = value[len(prefix):].split("-")
    if len(tail) < 2:
        return None
    return "-".join(tail[:-1]), tail[-1]


source_rows = []
for path in sorted((ROOT / "Vilion").glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=6, max_row=real_max_row(path), min_col=1, max_col=11, values_only=True),
        start=6,
    ):
        sku = values[10]
        if sku is None or str(sku).strip() == "":
            continue
        family = sku_family(str(sku))
        if not family:
            continue
        variation = normalized_variation(str(values[5]), family)
        tail = sku_tail(str(sku), family)
        if not tail or len(variation) < 2:
            continue
        source_rows.append({
            "file": path.name,
            "row": row_number,
            "product_id": str(values[0]),
            "family": family,
            "variation": variation,
            "sku": str(sku),
            "color_code": tail[0],
            "size_code": tail[1],
        })

# Learn the expected embedded codes from repeated source rows and reject isolated
# mismatches before building the lookup used for Seosm.
color_codes = defaultdict(Counter)
size_codes = defaultdict(Counter)
for item in source_rows:
    color_codes[(item["family"], item["variation"][0])][item["color_code"]] += 1
    size_codes[(item["family"], item["variation"][-1])][item["size_code"]] += 1

valid_source_rows = []
rejected_source_rows = []
for item in source_rows:
    expected_color = color_codes[(item["family"], item["variation"][0])].most_common(1)[0][0]
    expected_size = size_codes[(item["family"], item["variation"][-1])].most_common(1)[0][0]
    if item["color_code"] != expected_color or item["size_code"] != expected_size:
        rejected_source_rows.append(item)
    else:
        valid_source_rows.append(item)

source_index = defaultdict(list)
for item in valid_source_rows:
    source_index[(item["family"], item["variation"])].append(item)

updates = []
skips = []
file_counts = defaultdict(lambda: {"updates": 0, "skips": 0, "reasons": Counter(), "products": defaultdict(lambda: {"updates": 0, "skips": 0})})

for path in sorted((ROOT / "Seosm").glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=6, max_row=real_max_row(path), min_col=1, max_col=11, values_only=True),
        start=6,
    ):
        product_id = str(values[0])
        family = TARGET_FAMILIES.get(product_id)
        variation_text = str(values[5] or "")
        existing = values[10]
        if existing is not None and str(existing).strip() != "":
            reason = "target_sku_already_filled"
            matches = []
        elif family not in {name for _prefix, name in FAMILY_PREFIXES}:
            reason = "no_matching_vilion_sku_family"
            matches = []
        else:
            key = normalized_variation(variation_text, family, target=True)
            matches = source_index.get((family, key), [])
            normalized_skus = {item["sku"].lower() for item in matches}
            if len(normalized_skus) == 1:
                reason = ""
            elif not matches:
                reason = "no_exact_vilion_variant"
            else:
                reason = "ambiguous_vilion_sku"

        product_stats = file_counts[path.name]["products"][product_id]
        if not reason:
            sku_counts = Counter(item["sku"] for item in matches)
            sku = sku_counts.most_common(1)[0][0]
            source = next(item for item in matches if item["sku"].lower() == sku.lower())
            updates.append({
                "file": path.name,
                "row": row_number,
                "product_id": product_id,
                "product_name": str(values[2] or ""),
                "variation": variation_text,
                "family": family,
                "source_file": source["file"],
                "source_product_id": source["product_id"],
                "source_row": source["row"],
                "sku": sku,
            })
            file_counts[path.name]["updates"] += 1
            product_stats["updates"] += 1
        else:
            skips.append({
                "file": path.name,
                "row": row_number,
                "product_id": product_id,
                "product_name": str(values[2] or ""),
                "variation": variation_text,
                "family": family,
                "reason": reason,
                "candidates": sorted({item["sku"] for item in matches}),
            })
            file_counts[path.name]["skips"] += 1
            file_counts[path.name]["reasons"][reason] += 1
            product_stats["skips"] += 1

INSPECTION.mkdir(parents=True, exist_ok=True)
updates_path = INSPECTION / "proposed_seosm_sku_updates.csv"
with updates_path.open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(updates[0].keys()))
    writer.writeheader()
    writer.writerows(updates)

(INSPECTION / "skipped_seosm_skus.json").write_text(json.dumps(skips, ensure_ascii=False, indent=2), encoding="utf-8")
(INSPECTION / "rejected_vilion_source_skus.json").write_text(json.dumps(rejected_source_rows, ensure_ascii=False, indent=2), encoding="utf-8")

summary = {
    "total_updates": len(updates),
    "total_skips": len(skips),
    "rejected_source_rows": len(rejected_source_rows),
    "files": [],
}
for file_name in sorted(file_counts):
    stats = file_counts[file_name]
    summary["files"].append({
        "file": file_name,
        "updates": stats["updates"],
        "skips": stats["skips"],
        "skips_by_reason": dict(stats["reasons"]),
        "products": {pid: dict(values) for pid, values in stats["products"].items()},
    })
(INSPECTION / "proposed_seosm_sku_update_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
