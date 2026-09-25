import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok\Vilion")
WANTED = {"marun", "marun muda", "marun tua", "biru", "biru muda", "biru langit", "abu", "abu muda", "cream", "krem", "merah", "merah marun", "merah cabe", "terracota", "teracota"}


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


for path in sorted(ROOT.glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    seen = set()
    for row in sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=11, values_only=True):
        variation, sku = str(row[5] or ""), row[10]
        color = variation.split(",")[0].strip().lower()
        if color in WANTED and sku not in (None, ""):
            key = (str(row[0]), color, str(sku).rsplit("-", 1)[0])
            if key not in seen:
                seen.add(key)
                print(f"{row[0]} | {color} | {variation} -> {sku}")
