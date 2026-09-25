import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok\Vilion")


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


for path in sorted(ROOT.glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    products = defaultdict(list)
    names = {}
    for row in sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=11, values_only=True):
        pid = str(row[0])
        names[pid] = str(row[2])
        if row[10] not in (None, ""):
            products[pid].append((str(row[5]), str(row[10])))
    for pid, pairs in products.items():
        print(f"\n{pid} | {len(pairs)} | {names[pid]}")
        for variation, sku in pairs[:12]:
            print(f"  {variation} -> {sku}")
