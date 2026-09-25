import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok")
OUT = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok\inspection\workbook_inventory.json")


def real_max_row(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(value) for value in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


inventory = []
for store in ["Vilion", "Seosm"]:
    for path in sorted((ROOT / store).glob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["Template"]
        max_row = real_max_row(path)
        products = defaultdict(lambda: {"name": "", "rows": 0, "sku_filled": 0, "variations": []})
        filled = 0
        blank = 0
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=6, max_row=max_row, min_col=1, max_col=11, values_only=True),
            start=6,
        ):
            product_id, product_name, variation, sku = values[0], values[2], values[5], values[10]
            item = products[str(product_id)]
            item["name"] = str(product_name or "")
            item["rows"] += 1
            if sku is None or str(sku).strip() == "":
                blank += 1
            else:
                filled += 1
                item["sku_filled"] += 1
            if len(item["variations"]) < 8:
                item["variations"].append({"row": row_number, "variation": variation, "sku": sku})
        entry = {
            "store": store,
            "file": path.name,
            "path": str(path),
            "sheetnames": book.sheetnames,
            "data_rows": max_row - 5,
            "filled_skus": filled,
            "blank_skus": blank,
            "products": [{"product_id": key, **value} for key, value in products.items()],
        }
        inventory.append(entry)

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")

for entry in inventory:
    print(f"{entry['store']} | {entry['file']} | rows={entry['data_rows']} | filled={entry['filled_skus']} | blank={entry['blank_skus']}")
    for product in entry["products"]:
        print(f"  {product['product_id']} | rows={product['rows']} | sku={product['sku_filled']} | {product['name']}")
