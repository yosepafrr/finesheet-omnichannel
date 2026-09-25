import json
import re
import zipfile
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok")
OUT = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok\inspection\product_match_candidates.json")


def clean(value):
    text = str(value or "").lower()
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def tokens(value):
    return set(clean(value).split())


def jaccard(left, right):
    return len(left & right) / len(left | right) if left or right else 1.0


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


def load_products(store):
    products = {}
    for path in sorted((ROOT / store).glob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["Template"]
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=13, values_only=True),
            start=6,
        ):
            product_id = str(row[0])
            if product_id not in products:
                products[product_id] = {
                    "store": store,
                    "file": path.name,
                    "product_id": product_id,
                    "name": str(row[2] or ""),
                    "description": str(row[6] or ""),
                    "category": str(row[1] or ""),
                    "prices": set(),
                    "variations": [],
                    "rows": [],
                }
            product = products[product_id]
            product["prices"].add(str(row[8] or ""))
            product["variations"].append(str(row[5] or ""))
            product["rows"].append({"row": row_number, "variation": row[5], "sku": row[10]})
    return products


vilion = load_products("Vilion")
seosm = load_products("Seosm")
results = []
for target in seosm.values():
    candidates = []
    target_var = {clean(value) for value in target["variations"]}
    for source in vilion.values():
        source_var = {clean(value) for value in source["variations"]}
        variation_similarity = jaccard(target_var, source_var)
        name_similarity = jaccard(tokens(target["name"]), tokens(source["name"]))
        description_similarity = jaccard(tokens(target["description"]), tokens(source["description"]))
        sequence_similarity = SequenceMatcher(None, clean(target["description"]), clean(source["description"])).ratio()
        price_similarity = jaccard(target["prices"], source["prices"])
        score = 0.50 * variation_similarity + 0.18 * description_similarity + 0.12 * sequence_similarity + 0.12 * name_similarity + 0.08 * price_similarity
        candidates.append({
            "source_product_id": source["product_id"],
            "source_file": source["file"],
            "source_name": source["name"],
            "source_rows": len(source["rows"]),
            "source_filled_skus": sum(1 for row in source["rows"] if row["sku"] not in (None, "")),
            "score": round(score, 4),
            "variation_similarity": round(variation_similarity, 4),
            "variation_exact": target_var == source_var,
            "name_similarity": round(name_similarity, 4),
            "description_similarity": round(description_similarity, 4),
            "description_sequence": round(sequence_similarity, 4),
            "price_similarity": round(price_similarity, 4),
            "target_only_variations": sorted(target_var - source_var),
            "source_only_variations": sorted(source_var - target_var),
        })
    candidates.sort(key=lambda item: item["score"], reverse=True)
    results.append({
        "target_file": target["file"],
        "target_product_id": target["product_id"],
        "target_name": target["name"],
        "target_rows": len(target["rows"]),
        "target_variations": target["variations"],
        "top_candidates": candidates[:6],
    })

OUT.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
for result in results:
    print(f"\nTARGET {result['target_product_id']} | rows={result['target_rows']} | {result['target_name']}")
    for candidate in result["top_candidates"][:4]:
        print(
            f"  {candidate['source_product_id']} | rows={candidate['source_rows']} | sku={candidate['source_filled_skus']} | "
            f"score={candidate['score']} var={candidate['variation_similarity']} exact={candidate['variation_exact']} "
            f"desc={candidate['description_similarity']} seq={candidate['description_sequence']} name={candidate['name_similarity']} | "
            f"{candidate['source_name']}"
        )
