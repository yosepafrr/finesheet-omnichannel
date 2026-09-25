import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok")
TARGET_IDS = {
    "1734273961869411960",
    "1734274041964365432",
    "1729614622983883740",
    "1734274159432205944",
    "1729733053551053788",
    "1729733061894835164",
}


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


def clean_html(text):
    text = re.sub(r"<br\s*/?>|</p>|</li>", "\n", str(text or ""), flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


for store in ["Seosm", "Vilion"]:
    for path in sorted((ROOT / store).glob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["Template"]
        seen = set()
        for row in sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=13, values_only=True):
            pid = str(row[0])
            if pid not in TARGET_IDS or pid in seen:
                continue
            seen.add(pid)
            print(f"\n{store} {pid} | {row[2]} | price={row[8]} | weight={row[12]}")
            text = clean_html(row[6])
            print(text[:2500])
