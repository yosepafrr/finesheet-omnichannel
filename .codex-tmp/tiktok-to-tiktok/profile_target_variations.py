import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok\Seosm")


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


for path in sorted(ROOT.glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    groups = defaultdict(lambda: {"name": "", "values": []})
    for row in sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=11, values_only=True):
        pid = str(row[0])
        groups[pid]["name"] = str(row[2])
        groups[pid]["values"].append(str(row[5]))
    print(f"\nFILE {path.name}")
    for pid, group in groups.items():
        parts = [tuple(piece.strip() for piece in value.split(",")) for value in group["values"]]
        left = sorted({item[0] for item in parts})
        right = sorted({item[1] for item in parts if len(item) > 1})
        print(f"\n{pid} | {len(parts)} | {group['name']}")
        print(f"  first={left}")
        print(f"  second={right}")
        print(f"  values={group['values'][:12]}")
