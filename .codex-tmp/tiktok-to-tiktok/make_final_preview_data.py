import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok")
OUT = ROOT / "inspection" / "final_preview_rows.json"
ROWS = {
    1: [6, 11, 30, 42, 90, 96, 132, 168, 192, 228, 252, 258, 306, 354, 367],
    2: [6, 11, 12, 18, 24, 30, 36, 42, 47, 48, 54, 60, 66, 71],
    3: list(range(6, 18)),
    4: list(range(6, 18)),
}

result = []
for index in range(1, 5):
    path = ROOT / f"Tiktoksellercenter_batchedit_20260925_all_information_template_{index}_SKU_updated.xlsx"
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    rows = []
    for row_number in ROWS[index]:
        values = next(sheet.iter_rows(min_row=row_number, max_row=row_number, min_col=1, max_col=11, values_only=True))
        rows.append([row_number, str(values[0]), str(values[2]), str(values[5]), values[10]])
    result.append({"index": index, "rows": rows})

OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)

