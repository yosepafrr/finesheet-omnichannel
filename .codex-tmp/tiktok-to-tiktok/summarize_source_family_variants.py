import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok\Vilion")


def max_row(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return max(int(v) for v in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


def family(sku):
    value = str(sku).lower()
    for prefix, name in [
        ("3in1-p-", "3IN1-P"), ("jas-p-", "JAS-P"), ("rmp-p-", "RMP-P"),
        ("kmj-", "KMJ"), ("rc-p-", "RC-P"), ("cln-p-", "CLN-P"),
        ("jas-w-", "JAS-W"), ("3in1-a-", "3IN1-A"), ("dsi-", "DSI"),
    ]:
        if value.startswith(prefix):
            return name
    return "OTHER"


groups = defaultdict(list)
for path in sorted(ROOT.glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    sheet = book["Template"]
    for row in sheet.iter_rows(min_row=6, max_row=max_row(path), min_col=1, max_col=11, values_only=True):
        sku = row[10]
        if sku in (None, ""):
            continue
        groups[family(sku)].append((str(row[5]), str(sku)))

for name, pairs in sorted(groups.items()):
    colors = sorted({pair[0].split(",")[0].strip() for pair in pairs if "," in pair[0]})
    sizes = sorted({pair[0].split(",", 1)[1].strip() for pair in pairs if "," in pair[0]})
    print(f"{name}: rows={len(pairs)} colors={colors} sizes={sizes}")
