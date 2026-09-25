import csv
import hashlib
import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIR = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Tiktok to tiktok\Seosm")
UPDATES_PATH = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok\inspection\proposed_seosm_sku_updates.csv")
OUTPUT_DIR = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\tiktok-to-tiktok")
REPORT_PATH = OUTPUT_DIR / "verification_report.json"
SHEET_XML = "xl/worksheets/sheet1.xml"


def cell_pattern(coordinate: str) -> re.Pattern[str]:
    return re.compile(
        rf'<c\b(?=[^>]*\br="{re.escape(coordinate)}")[^>]*(?:/>|>.*?</c>)',
        flags=re.DOTALL,
    )


def real_max_row(xml: str) -> int:
    return max(int(value) for value in re.findall(r'<row\b[^>]*\br="(\d+)"', xml))


updates_by_file = defaultdict(list)
with UPDATES_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
    for row in csv.DictReader(handle):
        updates_by_file[row["file"]].append(row)

results = []
total_updates = 0
total_blanks = 0
for source_path in sorted(SOURCE_DIR.glob("*.xlsx")):
    rows = sorted(updates_by_file.get(source_path.name, []), key=lambda item: int(item["row"]))
    approved = {int(item["row"]): item for item in rows}
    output_path = OUTPUT_DIR / f"{source_path.stem}_SKU_updated.xlsx"
    if not output_path.exists():
        raise RuntimeError(f"Missing output: {output_path}")

    with zipfile.ZipFile(source_path, "r") as before_zip, zipfile.ZipFile(output_path, "r") as after_zip:
        assert before_zip.namelist() == after_zip.namelist()
        source_xml = before_zip.read(SHEET_XML).decode("utf-8")
        data_max_row = real_max_row(source_xml)
        if not rows:
            assert source_path.read_bytes() == output_path.read_bytes(), f"Zero-update file changed: {source_path.name}"
            unchanged_entries = len(before_zip.namelist())
        else:
            for name in before_zip.namelist():
                if name != SHEET_XML:
                    assert before_zip.read(name) == after_zip.read(name), f"Unexpected package change: {source_path.name}: {name}"
            restored_xml = after_zip.read(SHEET_XML).decode("utf-8")
            for update in rows:
                coordinate = f"K{update['row']}"
                pattern = cell_pattern(coordinate)
                source_match = pattern.search(source_xml)
                output_match = pattern.search(restored_xml)
                assert source_match and output_match, (source_path.name, coordinate)
                assert f"<t>{update['sku']}</t>" in output_match.group(0), (source_path.name, coordinate)
                restored_xml = restored_xml[: output_match.start()] + source_match.group(0) + restored_xml[output_match.end() :]
            assert restored_xml == source_xml, f"Worksheet changed outside approved cells: {source_path.name}"
            unchanged_entries = len(before_zip.namelist()) - 1

    before_book = load_workbook(source_path, read_only=True, data_only=False)
    after_book = load_workbook(output_path, read_only=True, data_only=False)
    assert before_book.sheetnames == after_book.sheetnames
    before_sheet = before_book["Template"]
    after_sheet = after_book["Template"]
    before_rows = before_sheet.iter_rows(min_row=6, max_row=data_max_row, min_col=1, max_col=11, values_only=True)
    after_rows = after_sheet.iter_rows(min_row=6, max_row=data_max_row, min_col=1, max_col=11, values_only=True)
    blanks = 0
    verified = set()
    for row_number, (before_values, after_values) in enumerate(zip(before_rows, after_rows), start=6):
        before_sku = before_values[10]
        after_sku = after_values[10]
        if row_number in approved:
            update = approved[row_number]
            assert before_sku is None or str(before_sku).strip() == ""
            assert after_sku == update["sku"], (source_path.name, row_number, after_sku, update["sku"])
            assert str(before_values[0]) == update["product_id"]
            assert str(before_values[5]) == update["variation"]
            verified.add(row_number)
        else:
            assert after_sku == before_sku, (source_path.name, row_number, before_sku, after_sku)
        if after_sku is None or str(after_sku).strip() == "":
            blanks += 1
    assert verified == set(approved)

    results.append({
        "file": output_path.name,
        "verified_updates": len(rows),
        "remaining_blank_skus": blanks,
        "unchanged_package_entries": unchanged_entries,
        "sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "size_bytes": output_path.stat().st_size,
    })
    total_updates += len(rows)
    total_blanks += blanks

report = {
    "verified_total_updates": total_updates,
    "verified_total_remaining_blank_skus": total_blanks,
    "unexpected_value_diffs": 0,
    "unexpected_package_changes": 0,
    "files": results,
}
REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
