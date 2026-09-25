import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


source_dir = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs")
template_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")
output_dir = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\inspection")
output_dir.mkdir(parents=True, exist_ok=True)

product_rules = {
    "26336754672": ("RC", "Anak"),
    "23422488757": ("JAS", "Anak"),
    "22656255576": ("3IN1", "Pria"),
    "22375612618": ("JC", "Wanita"),
    "22226266149": ("JC", "Pria"),
    "21876818960": ("JC", "Wanita"),
    "18735176975": ("RMP", "Pria"),
    "18653316238": ("JAS", "Wanita"),
    "18584652576": ("JC", "Pria"),
    "16092916076": ("JR", "Pria"),
    "15991304773": ("JAS", "Pria"),
    "14113818693": ("CLN", "Pria"),
    "11397737616": ("3IN1", "Pria"),
}

color_aliases = {
    "abu": "Abu",
    "abu abu": "Abu",
    "baby pink": "Baby Pink",
    "biru": "Biru BCA",
    "biru bca": "Biru BCA",
    "biru langit": "Biru Langit",
    "sky blue": "Biru Langit",
    "iced blue": "Biru Langit",
    "burgundy": "Marun Tua",
    "coklat": "Coklat Mahogani",
    "coklat mahogani": "Coklat Mahogani",
    "cream": "Krem",
    "hitam": "Hitam",
    "krem": "Krem",
    "krem cream": "Krem",
    "mahogani": "Coklat Mahogani",
    "maroon": "Marun Muda",
    "marun": "Marun Muda",
    "marun muda": "Marun Muda",
    "marun tua": "Marun Tua",
    "merah cabe": "Merah Cabe",
    "merah cabai": "Merah Cabe",
    "merah maroon": "Marun Muda",
    "navy": "Navy",
    "navy blue": "Navy",
    "pink fanta": "Pink Fanta",
    "putih": "Putih",
    "sage": "Sage",
    "sage hijau toska": "Sage",
    "terracota": "Terracota",
}


def normalize_words(value):
    text = str(value or "").strip().casefold()
    text = re.sub(r"[^0-9a-z]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_color(value):
    return color_aliases.get(normalize_words(value))


def normalize_size(value):
    text = str(value or "").strip()
    text = re.sub(r"\bThn\b", "Tahun", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    if re.fullmatch(r"[0-9]+", text):
        return text
    return text.upper().replace(" TAHUN", " Tahun")


master_by_key = {}
master_conflicts = defaultdict(set)
master_files = [
    "Master_Produk_Anak_SKU_Revisi.xlsx",
    "Master_Produk_Wanita.xlsx",
    "SKU_PRODUK_PRIA_INCLUDE_CELANA_UPDATED.xlsx",
]
for name in master_files:
    workbook = load_workbook(source_dir / name, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        sku, _product_name, _category, color, size, gender = row[:6]
        if not sku:
            continue
        sku_text = str(sku).strip()
        prefix = sku_text.split("-")[0].upper()
        key = (prefix, str(gender).strip(), str(color).strip(), normalize_size(size))
        master_conflicts[key].add(sku_text)
        master_by_key[key] = sku_text
    workbook.close()

conflicts = {key: values for key, values in master_conflicts.items() if len(values) > 1}
if conflicts:
    raise RuntimeError(f"Ambiguous master SKU keys: {conflicts}")

workbook = load_workbook(template_path, read_only=False, data_only=False)
sheet = workbook[workbook.sheetnames[0]]
updates = []
skipped = []
for row_number in range(7, sheet.max_row + 1):
    product_id = str(sheet.cell(row_number, 1).value or "").strip()
    product_name = sheet.cell(row_number, 2).value
    variation_name = sheet.cell(row_number, 4).value
    current_sku = sheet.cell(row_number, 6).value
    if current_sku is not None and str(current_sku).strip() != "":
        continue
    rule = product_rules.get(product_id)
    if not rule:
        skipped.append({"row": row_number, "product_id": product_id, "product_name": product_name, "variation_name": variation_name, "reason": "product_not_mapped"})
        continue
    parts = [part.strip() for part in str(variation_name or "").split(",")]
    if len(parts) < 2:
        skipped.append({"row": row_number, "product_id": product_id, "product_name": product_name, "variation_name": variation_name, "reason": "variation_has_no_color_and_size"})
        continue
    color = normalize_color(parts[0])
    size = normalize_size(parts[-1])
    if not color:
        skipped.append({"row": row_number, "product_id": product_id, "product_name": product_name, "variation_name": variation_name, "reason": "color_not_in_master"})
        continue
    prefix, gender = rule
    key = (prefix, gender, color, size)
    sku = master_by_key.get(key)
    if not sku:
        skipped.append({"row": row_number, "product_id": product_id, "product_name": product_name, "variation_name": variation_name, "reason": "exact_master_combination_not_found", "lookup_key": key})
        continue
    updates.append({
        "row": row_number,
        "product_id": product_id,
        "product_name": product_name,
        "variation_name": variation_name,
        "master_color": color,
        "master_size": size,
        "sku": sku,
    })

with (output_dir / "proposed_sku_updates.csv").open("w", newline="", encoding="utf-8-sig") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(updates[0].keys()))
    writer.writeheader()
    writer.writerows(updates)
with (output_dir / "skipped_missing_skus.csv").open("w", newline="", encoding="utf-8-sig") as handle:
    fieldnames = ["row", "product_id", "product_name", "variation_name", "reason", "lookup_key"]
    writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(skipped)

summary = {
    "proposed_updates": len(updates),
    "skipped_empty_skus": len(skipped),
    "updates_by_product": dict(Counter(f'{item["product_id"]} | {item["product_name"]}' for item in updates)),
    "skips_by_reason": dict(Counter(item["reason"] for item in skipped)),
}
(output_dir / "proposed_sku_update_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
