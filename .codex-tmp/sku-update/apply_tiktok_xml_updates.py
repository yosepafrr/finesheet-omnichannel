import csv
import html
import re
import shutil
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIR = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Shopee to tiktok\Tiktoksellercenter_batchedit_20260925_all_information_template")
UPDATES_PATH = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok\inspection\proposed_tiktok_sku_updates.csv")
OUTPUT_DIR = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok")
SHEET_XML = "xl/worksheets/sheet1.xml"


def cell_pattern(coordinate: str) -> re.Pattern[str]:
    return re.compile(
        rf'<c\b(?=[^>]*\br="{re.escape(coordinate)}")[^>]*(?:/>|>.*?</c>)',
        flags=re.DOTALL,
    )


updates_by_file: dict[str, list[dict[str, str]]] = defaultdict(list)
with UPDATES_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
    for row in csv.DictReader(handle):
        updates_by_file[row["file"]].append(row)

source_files = sorted(SOURCE_DIR.glob("*.xlsx"))
if len(source_files) != 5:
    raise RuntimeError(f"Expected 5 TikTok workbooks, found {len(source_files)}")

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
for source_path in source_files:
    rows = sorted(updates_by_file.get(source_path.name, []), key=lambda item: int(item["row"]))
    output_path = OUTPUT_DIR / f"{source_path.stem}_SKU_updated.xlsx"

    if not rows:
        shutil.copy2(source_path, output_path)
        print(f"Copied unchanged (0 safe matches): {output_path}")
        continue

    # TikTok stores blank cells as a shared-string reference whose resolved value is
    # empty. Validate the resolved values before replacing the underlying XML cells.
    source_book = load_workbook(source_path, read_only=True, data_only=False)
    source_sheet = source_book["Template"]
    for update in rows:
        row_number = int(update["row"])
        existing = source_sheet.cell(row_number, 11).value
        if existing is not None and str(existing).strip() != "":
            raise RuntimeError(f"Refusing to overwrite populated cell K{row_number} in {source_path.name}")
        if str(source_sheet.cell(row_number, 1).value) != update["product_id"]:
            raise RuntimeError(f"Product mismatch at row {row_number} in {source_path.name}")
        if str(source_sheet.cell(row_number, 6).value) != update["variation"]:
            raise RuntimeError(f"Variation mismatch at row {row_number} in {source_path.name}")

    with zipfile.ZipFile(source_path, "r") as source_zip:
        sheet_xml = source_zip.read(SHEET_XML).decode("utf-8")
        original_sheet_xml = sheet_xml

        for update in rows:
            row_number = int(update["row"])
            coordinate = f"K{row_number}"
            pattern = cell_pattern(coordinate)
            match = pattern.search(sheet_xml)
            if not match:
                raise RuntimeError(f"Cell {coordinate} not found in {source_path.name}")
            source_cell = match.group(0)
            opening_match = re.match(r"<c\b[^>]*", source_cell)
            if not opening_match:
                raise RuntimeError(f"Could not parse {coordinate} in {source_path.name}")
            opening_tag = re.sub(r'\s+t="[^"]*"', "", opening_match.group(0).rstrip("/"))
            opening_tag += ' t="inlineStr">'
            sku = html.escape(update["sku"], quote=False)
            replacement = f"{opening_tag}<is><t>{sku}</t></is></c>"
            sheet_xml = sheet_xml[: match.start()] + replacement + sheet_xml[match.end() :]

        if sheet_xml == original_sheet_xml:
            raise RuntimeError(f"No changes made for {source_path.name}")

        with zipfile.ZipFile(output_path, "w") as output_zip:
            for item in source_zip.infolist():
                data = sheet_xml.encode("utf-8") if item.filename == SHEET_XML else source_zip.read(item.filename)
                output_zip.writestr(item, data)

    print(f"Wrote {len(rows)} SKU values: {output_path}")
