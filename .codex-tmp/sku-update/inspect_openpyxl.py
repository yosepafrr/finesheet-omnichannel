import json
from pathlib import Path

from openpyxl import load_workbook


source_dir = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs")
output_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\inspection\openpyxl_summary.json")
output_path.parent.mkdir(parents=True, exist_ok=True)

summary = []
files = sorted(source_dir.glob("*.xlsx"))
sanitized_template = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")
for original_path in files:
    file_path = sanitized_template if original_path.name.startswith("mass_update_sales") else original_path
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        samples = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), max_col=min(sheet.max_column, 50), values_only=True):
            samples.append(list(row))
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "sample_values": samples,
            }
        )
    workbook.close()
    summary.append({"file": original_path.name, "opened_from": str(file_path), "sheets": sheets})

output_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
