import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


source_dir = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs")
files = [
    source_dir / "Master_Produk_Anak_SKU_Revisi.xlsx",
    source_dir / "Master_Produk_Wanita.xlsx",
    source_dir / "SKU_PRODUK_PRIA_INCLUDE_CELANA_UPDATED.xlsx",
]

for file_path in files:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    print(f"FILE: {file_path.name}; ROWS: {len(rows)}")
    for label, index in [("kategori", 2), ("warna", 3), ("size", 4), ("gender", 5)]:
        values = Counter(str(row[index]) for row in rows if row[index] not in (None, ""))
        print(f"  {label}: {json.dumps(dict(values), ensure_ascii=False, sort_keys=True)}")
    print("  last samples:")
    for row in rows[-12:]:
        print("   ", json.dumps(row[:6], ensure_ascii=False, default=str))
    workbook.close()
