import io
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


source_dir = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Shopee to tiktok")
output_dir = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok\inspection")
output_dir.mkdir(parents=True, exist_ok=True)


def workbook_bytes_with_valid_panes(file_path):
    source_bytes = file_path.read_bytes()
    input_buffer = io.BytesIO(source_bytes)
    output_buffer = io.BytesIO()
    replacements = {
        b'activePane="bottom_left"': b'activePane="bottomLeft"',
        b'activePane="top_left"': b'activePane="topLeft"',
        b'activePane="bottom_right"': b'activePane="bottomRight"',
        b'activePane="top_right"': b'activePane="topRight"',
    }
    changed = False
    with zipfile.ZipFile(input_buffer, "r") as source_zip, zipfile.ZipFile(output_buffer, "w") as destination_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                for old, new in replacements.items():
                    if old in data:
                        data = data.replace(old, new)
                        changed = True
            destination_zip.writestr(item, data)
    return output_buffer.getvalue() if changed else source_bytes, changed


summary = []
for file_path in sorted(source_dir.rglob("*.xlsx")):
    workbook_bytes, pane_sanitized = workbook_bytes_with_valid_panes(file_path)
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=False, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        top_rows = []
        for row_number in range(1, min(sheet.max_row, 20) + 1):
            populated = []
            for column_number in range(1, min(sheet.max_column, 250) + 1):
                value = sheet.cell(row_number, column_number).value
                if value not in (None, ""):
                    populated.append({
                        "cell": f"{get_column_letter(column_number)}{row_number}",
                        "value": str(value)[:500],
                    })
            if populated:
                top_rows.append({"row": row_number, "populated": populated})
        sheets.append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "top_rows": top_rows,
        })
    workbook.close()
    summary.append({
        "file": file_path.name,
        "relative_path": str(file_path.relative_to(source_dir)),
        "pane_sanitized_in_memory": pane_sanitized,
        "sheets": sheets,
    })

(output_dir / "workbook_structure.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(summary, ensure_ascii=False, indent=2))
