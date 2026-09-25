from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


OUTPUT_DIR = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok")
QA_PATH = OUTPUT_DIR / "inspection" / "tiktok-final-visual-qa.xlsx"
RANGES = {1: (6, 20), 2: (6, 15), 3: (6, 7), 4: (6, 20), 5: (6, 20)}

qa_book = Workbook()
qa_book.remove(qa_book.active)

for index, (start_row, end_row) in RANGES.items():
    source_path = OUTPUT_DIR / f"Tiktoksellercenter_batchedit_20260925_all_information_template_{index}_SKU_updated.xlsx"
    source_book = load_workbook(source_path, read_only=True, data_only=False)
    source_sheet = source_book["Template"]
    qa_sheet = qa_book.create_sheet(f"File {index}")
    qa_sheet.append(["Baris sumber", "ID Produk", "Nilai Variasi", "SKU Penjual"])
    for row_number, values in enumerate(
        source_sheet.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=11, values_only=True),
        start=start_row,
    ):
        qa_sheet.append([row_number, values[0], values[5], values[10]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in qa_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    qa_sheet.freeze_panes = "A2"
    qa_sheet.auto_filter.ref = qa_sheet.dimensions
    qa_sheet.column_dimensions["A"].width = 14
    qa_sheet.column_dimensions["B"].width = 24
    qa_sheet.column_dimensions["C"].width = 28
    qa_sheet.column_dimensions["D"].width = 28

QA_PATH.parent.mkdir(parents=True, exist_ok=True)
qa_book.save(QA_PATH)
print(QA_PATH)

