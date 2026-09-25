import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


source_path = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\mass_update_sales_info_615078427_20260925103824.xlsx")
output_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\mass_update_sales_info_615078427_20260925103824_SKU_updated.xlsx")
updates_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\inspection\proposed_sku_updates.csv")
sanitized_source_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")
sanitized_output_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_updated_sanitized_for_verification.xlsx")

with updates_path.open("r", encoding="utf-8-sig", newline="") as handle:
    updates = {int(row["row"]): row["sku"] for row in csv.DictReader(handle)}

with zipfile.ZipFile(source_path, "r") as source_zip, zipfile.ZipFile(output_path, "r") as output_zip:
    assert source_zip.namelist() == output_zip.namelist()
    unchanged_entries = []
    for name in source_zip.namelist():
        if name == "xl/worksheets/sheet1.xml":
            continue
        assert source_zip.read(name) == output_zip.read(name), f"Unexpected package change: {name}"
        unchanged_entries.append(name)

    source_xml = source_zip.read("xl/worksheets/sheet1.xml").decode("utf-8")
    output_xml = output_zip.read("xl/worksheets/sheet1.xml").decode("utf-8")
    restored_xml = output_xml
    for row_number, sku in sorted(updates.items()):
        coordinate = f"F{row_number}"
        cell_pattern = re.compile(
            rf'<c\b(?=[^>]*\br="{re.escape(coordinate)}")[^>]*>.*?</c>',
            flags=re.DOTALL,
        )
        source_match = cell_pattern.search(source_xml)
        output_match = cell_pattern.search(restored_xml)
        assert source_match and output_match, coordinate
        assert not re.search(r"<(?:v|is)\b", source_match.group(0)), f"Source cell was not blank: {coordinate}"
        assert 't="inlineStr"' in output_match.group(0), coordinate
        assert f"<t>{sku}</t>" in output_match.group(0), f"Wrong SKU in {coordinate}"
        restored_xml = restored_xml[: output_match.start()] + source_match.group(0) + restored_xml[output_match.end() :]
    assert restored_xml == source_xml, "Worksheet contains changes outside approved SKU cells"

with zipfile.ZipFile(output_path, "r") as source_zip, zipfile.ZipFile(sanitized_output_path, "w") as destination_zip:
    for item in source_zip.infolist():
        data = source_zip.read(item.filename)
        if item.filename == "xl/worksheets/sheet1.xml":
            data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
        destination_zip.writestr(item, data)

before_book = load_workbook(sanitized_source_path, read_only=False, data_only=False)
after_book = load_workbook(sanitized_output_path, read_only=False, data_only=False)
assert before_book.sheetnames == after_book.sheetnames

value_diffs = []
style_diffs = []
formula_count = 0
formula_errors = []
error_values = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
for before_sheet, after_sheet in zip(before_book.worksheets, after_book.worksheets):
    assert before_sheet.title == after_sheet.title
    assert before_sheet.max_row == after_sheet.max_row
    assert before_sheet.max_column == after_sheet.max_column
    assert before_sheet.freeze_panes == after_sheet.freeze_panes
    assert set(before_sheet.merged_cells.ranges) == set(after_sheet.merged_cells.ranges)
    for row in range(1, before_sheet.max_row + 1):
        for column in range(1, before_sheet.max_column + 1):
            before_cell = before_sheet.cell(row, column)
            after_cell = after_sheet.cell(row, column)
            if before_cell.value != after_cell.value:
                value_diffs.append((after_sheet.title, after_cell.coordinate, before_cell.value, after_cell.value))
            if before_cell._style != after_cell._style:
                style_diffs.append((after_sheet.title, after_cell.coordinate))
            if isinstance(after_cell.value, str) and after_cell.value.startswith("="):
                formula_count += 1
            if after_cell.value in error_values:
                formula_errors.append((after_sheet.title, after_cell.coordinate, after_cell.value))

allowed_diffs = {("Sheet1", f"F{row}") for row in updates}
assert {(sheet, cell) for sheet, cell, _before, _after in value_diffs} == allowed_diffs
assert not style_diffs, style_diffs[:10]
assert not formula_errors, formula_errors

preserved_existing = 0
remaining_blank = 0
sheet_before = before_book["Sheet1"]
sheet_after = after_book["Sheet1"]
for row in range(7, sheet_after.max_row + 1):
    before_value = sheet_before.cell(row, 6).value
    after_value = sheet_after.cell(row, 6).value
    if before_value is not None and str(before_value).strip() != "":
        assert after_value == before_value
        preserved_existing += 1
    if after_value is None or str(after_value).strip() == "":
        remaining_blank += 1

result = {
    "verified_updates": len(updates),
    "unchanged_package_entries": len(unchanged_entries),
    "unexpected_value_diffs": 0,
    "style_diffs": len(style_diffs),
    "preserved_existing_skus": preserved_existing,
    "remaining_blank_skus": remaining_blank,
    "formula_count": formula_count,
    "formula_errors": len(formula_errors),
    "output_size_bytes": output_path.stat().st_size,
    "sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
}
print(json.dumps(result, ensure_ascii=False, indent=2))
