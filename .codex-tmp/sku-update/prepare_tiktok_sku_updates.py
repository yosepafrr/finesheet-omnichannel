import csv
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


root = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Shopee to tiktok")
source_path = next(root.glob("mass_update_sales*.xlsx"))
tiktok_paths = sorted(root.rglob("Tiktoksellercenter*.xlsx"))
output_dir = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok\inspection")
output_dir.mkdir(parents=True, exist_ok=True)

# TikTok product ID -> exact Shopee listing, expected SKU family, optional fixed color.
product_mapping = {
    "1729441743218904028": ("11397737616", "3IN1-P", None),
    "1729441757499001820": ("15894005218", "JAS-P", None),
    "1729515703285615580": ("18735176975", "RMP-P", None),
    "1729601659603553244": ("16123167377", "RMP-P", None),
    "1729605595651213276": ("17696424417", "KMJ", None),
    "1729614622983883740": ("22772266116", "RC-P", None),
    "1729670441825896412": ("14113818693", "CLN-P", None),
    "1729733026619755484": ("13276749795", "JAS-P", None),
    "1729733033814035420": ("22656255576", "3IN1-P", None),
    "1729733049820154844": ("22656255576", "3IN1-P", "hitam"),
    "1729733072589065180": ("11397737616", "3IN1-P", None),
    "1729733090698234844": ("22656255576", "3IN1-P", None),
    "1729733053551053788": ("23422488757", "JAS-A", None),
    "1729589401601869788": ("47160852979", "DSI-KPKP", None),
    "1729617524176488412": ("20054716390", "DSI-PNJNG", None),
    "1729733061894835164": ("24030445594", "3IN1-A", None),
    "1729441747251923932": ("18437126015", "JAS-W", None),
    "1729515705665948636": ("18437126015", "JAS-W", None),
}


def sanitized_bytes(file_path):
    input_buffer = io.BytesIO(file_path.read_bytes())
    output_buffer = io.BytesIO()
    replacements = {
        b'activePane="bottom_left"': b'activePane="bottomLeft"',
        b'activePane="top_left"': b'activePane="topLeft"',
        b'activePane="bottom_right"': b'activePane="bottomRight"',
        b'activePane="top_right"': b'activePane="topRight"',
    }
    changed = False
    with zipfile.ZipFile(input_buffer, "r") as source_zip, zipfile.ZipFile(output_buffer, "w") as destination_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                for old, new in replacements.items():
                    if old in data:
                        data = data.replace(old, new)
                        changed = True
            destination_zip.writestr(item, data)
    return output_buffer.getvalue() if changed else file_path.read_bytes()


def norm_text(value):
    text = str(value or "").casefold().strip()
    text = re.sub(r"[^0-9a-z]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


color_aliases = {
    "abu abu": "abu", "abu": "abu", "baby pink": "baby pink", "biru": "biru bca",
    "biru bca": "biru bca", "biru langit": "biru langit", "sky blue": "biru langit",
    "iced blue": "biru langit", "burgundy": "marun tua", "coklat": "coklat mahogani",
    "coklat mahogani": "coklat mahogani", "cream": "krem", "krem": "krem",
    "hitam": "hitam", "mahogani": "coklat mahogani", "maroon": "marun muda",
    "marun": "marun muda", "marun muda": "marun muda", "marun tua": "marun tua",
    "merah cabe": "merah cabe", "merah cabai": "merah cabe", "merah maroon": "marun muda",
    "navy": "navy", "navy blue": "navy", "pink fanta": "pink fanta", "putih": "putih",
    "sage": "sage", "sage hijau toska": "sage", "terracota": "terracota",
}


def norm_component(value, is_color=False):
    token = norm_text(value)
    if is_color:
        token = color_aliases.get(token, token)
    token = re.sub(r"\b(thn|t)\b", "tahun", token)
    if token == "3xl":
        token = "xxxl"
    return token


def norm_variation(value, fixed_color=None):
    raw = str(value or "").strip()
    if fixed_color and "," not in raw:
        return f"{fixed_color}|{norm_component(raw)}"
    if not raw:
        return ""
    parts = [part.strip() for part in raw.split(",")]
    return "|".join(norm_component(part, is_color=(index == 0)) for index, part in enumerate(parts))


def parse_sku(sku):
    parts = str(sku).strip().upper().split("-")
    if len(parts) >= 4 and parts[1] in {"P", "W", "A"}:
        return f"{parts[0]}-{parts[1]}", parts[2], "-".join(parts[3:])
    if len(parts) >= 3 and parts[0] == "KMJ":
        return "KMJ", parts[1], "-".join(parts[2:])
    return None, None, None


source_book = load_workbook(io.BytesIO(sanitized_bytes(source_path)), read_only=False, data_only=True)
source_sheet = source_book[source_book.sheetnames[0]]
source_groups = defaultdict(lambda: defaultdict(set))
all_source_rows = []
for row in range(7, source_sheet.max_row + 1):
    product_id = str(source_sheet.cell(row, 1).value or "").strip()
    variation = source_sheet.cell(row, 4).value
    parent_sku = source_sheet.cell(row, 5).value
    variation_sku = source_sheet.cell(row, 6).value
    sku = variation_sku if variation_sku is not None and str(variation_sku).strip() else parent_sku
    if not product_id or sku is None or str(sku).strip() == "":
        continue
    sku = str(sku).strip()
    key = norm_variation(variation)
    source_groups[product_id][key].add(sku)
    all_source_rows.append({"row": row, "product_id": product_id, "variation_key": key, "sku": sku})
source_book.close()

# Learn the dominant embedded color and size codes from all parseable source SKUs.
color_code_counts = defaultdict(Counter)
size_code_counts = defaultdict(Counter)
for item in all_source_rows:
    parts = item["variation_key"].split("|")
    family, color_code, size_code = parse_sku(item["sku"])
    if family and len(parts) >= 2:
        color_code_counts[(family, parts[0])][color_code] += 1
        size_code_counts[(family, parts[-1])][size_code] += 1
dominant_color_code = {key: counts.most_common(1)[0][0] for key, counts in color_code_counts.items()}
dominant_size_code = {key: counts.most_common(1)[0][0] for key, counts in size_code_counts.items()}


def semantic_check(sku, expected_family, variation_key):
    if expected_family.startswith("DSI-"):
        expected_sku = "dsi-kpkp" if expected_family == "DSI-KPKP" else "dsi-pnjng"
        return sku.casefold() == expected_sku, {"expected_sku": expected_sku}
    family, color_code, size_code = parse_sku(sku)
    variation_parts = variation_key.split("|")
    expected_color_code = dominant_color_code.get((expected_family, variation_parts[0])) if len(variation_parts) >= 2 else None
    expected_size_code = dominant_size_code.get((expected_family, variation_parts[-1])) if len(variation_parts) >= 2 else None
    valid = (
        family == expected_family
        and color_code == expected_color_code
        and size_code == expected_size_code
    )
    return valid, {
        "parsed_family": family,
        "expected_family": expected_family,
        "parsed_color_code": color_code,
        "expected_color_code": expected_color_code,
        "parsed_size_code": size_code,
        "expected_size_code": expected_size_code,
    }


family_lookup = defaultdict(lambda: defaultdict(set))
family_sources = defaultdict(lambda: defaultdict(lambda: defaultdict(set)))
for item in all_source_rows:
    family, _color_code, _size_code = parse_sku(item["sku"])
    if not family or item["sku"] == "0":
        continue
    valid, _details = semantic_check(item["sku"], family, item["variation_key"])
    if valid:
        family_lookup[family][item["variation_key"]].add(item["sku"])
        family_sources[family][item["variation_key"]][item["sku"]].add(item["product_id"])

all_updates = []
all_skips = []
file_reports = []
for tiktok_path in tiktok_paths:
    workbook = load_workbook(io.BytesIO(sanitized_bytes(tiktok_path)), read_only=False, data_only=True)
    sheet = workbook["Template"]
    updates = []
    skips = []
    for row in range(6, sheet.max_row + 1):
        product_id = str(sheet.cell(row, 1).value or "").strip()
        product_name = sheet.cell(row, 3).value
        variation = sheet.cell(row, 6).value
        existing_sku = sheet.cell(row, 11).value
        if not product_id:
            continue
        if existing_sku is not None and str(existing_sku).strip() != "":
            skips.append({"row": row, "product_id": product_id, "variation": variation, "reason": "seller_sku_already_filled"})
            continue
        mapping = product_mapping.get(product_id)
        if not mapping:
            skips.append({"row": row, "product_id": product_id, "variation": variation, "reason": "product_not_mapped"})
            continue
        source_product_id, expected_family, fixed_color = mapping
        target_key = norm_variation(variation, fixed_color=fixed_color)
        source_key = "" if expected_family.startswith("DSI-") and norm_text(variation) == "default" else target_key
        direct_candidates = source_groups[source_product_id].get(source_key, set())
        valid_direct = set()
        invalid_direct = []
        for candidate in direct_candidates:
            valid, details = semantic_check(candidate, expected_family, source_key)
            if candidate != "0" and valid:
                valid_direct.add(candidate)
            else:
                invalid_direct.append({"sku": candidate, **details})

        source_basis = "exact_product"
        source_ids_used = [source_product_id]
        if len(valid_direct) == 1:
            sku = next(iter(valid_direct))
        elif not expected_family.startswith("DSI-"):
            family_candidates = family_lookup[expected_family].get(source_key, set())
            if len(family_candidates) == 1:
                sku = next(iter(family_candidates))
                source_basis = "validated_family_variant"
                source_ids_used = sorted(family_sources[expected_family][source_key][sku])
            else:
                reason = "source_sku_failed_semantic_check" if direct_candidates and not valid_direct else "no_unique_valid_source_variant"
                skips.append({
                    "row": row, "product_id": product_id, "product_name": product_name, "variation": variation,
                    "source_product_id": source_product_id, "reason": reason, "source_key": source_key,
                    "direct_candidates": sorted(direct_candidates), "invalid_direct": invalid_direct,
                    "family_candidates": sorted(family_candidates),
                })
                continue
        else:
            reason = "source_sku_failed_semantic_check" if direct_candidates else "no_unique_valid_source_variant"
            skips.append({
                "row": row, "product_id": product_id, "product_name": product_name, "variation": variation,
                "source_product_id": source_product_id, "reason": reason, "source_key": source_key,
                "direct_candidates": sorted(direct_candidates), "invalid_direct": invalid_direct,
            })
            continue
        update = {
            "file": tiktok_path.name,
            "row": row,
            "product_id": product_id,
            "product_name": product_name,
            "variation": variation,
            "source_product_id": ",".join(source_ids_used),
            "source_basis": source_basis,
            "sku": sku,
        }
        updates.append(update)
        all_updates.append(update)
    workbook.close()
    for skip in skips:
        skip["file"] = tiktok_path.name
        all_skips.append(skip)
    file_reports.append({
        "file": tiktok_path.name,
        "updates": len(updates),
        "skips": len(skips),
        "skips_by_reason": dict(Counter(item["reason"] for item in skips)),
        "updates_by_product": dict(Counter(item["product_id"] for item in updates)),
    })

with (output_dir / "proposed_tiktok_sku_updates.csv").open("w", newline="", encoding="utf-8-sig") as handle:
    fields = ["file", "row", "product_id", "product_name", "variation", "source_product_id", "source_basis", "sku"]
    writer = csv.DictWriter(handle, fieldnames=fields)
    writer.writeheader()
    writer.writerows(all_updates)
with (output_dir / "skipped_tiktok_skus.json").open("w", encoding="utf-8") as handle:
    json.dump(all_skips, handle, ensure_ascii=False, indent=2)

summary = {
    "total_updates": len(all_updates),
    "total_skips": len(all_skips),
    "files": file_reports,
    "global_skips_by_reason": dict(Counter(item["reason"] for item in all_skips)),
}
(output_dir / "proposed_tiktok_sku_update_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
