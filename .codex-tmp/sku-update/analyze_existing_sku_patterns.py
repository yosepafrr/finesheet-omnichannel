import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


source_dir = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs")
template_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")

master_by_sku = {}
for name in ["Master_Produk_Anak_SKU_Revisi.xlsx", "Master_Produk_Wanita.xlsx", "SKU_PRODUK_PRIA_INCLUDE_CELANA_UPDATED.xlsx"]:
    workbook = load_workbook(source_dir / name, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:
            master_by_sku[str(row[0]).strip().casefold()] = {
                "sku": str(row[0]).strip(),
                "name": row[1],
                "category": row[2],
                "color": row[3],
                "size": str(row[4]) if row[4] is not None else None,
                "gender": row[5],
            }
    workbook.close()

workbook = load_workbook(template_path, read_only=False, data_only=True)
sheet = workbook[workbook.sheetnames[0]]
color_map = defaultdict(Counter)
size_map = defaultdict(Counter)
product_map = defaultdict(Counter)
matched = 0
unmatched = Counter()
for row_number in range(7, sheet.max_row + 1):
    product_name = sheet.cell(row_number, 2).value
    variation_name = sheet.cell(row_number, 4).value
    sku = sheet.cell(row_number, 6).value
    if sku is None or str(sku).strip() == "":
        continue
    master = master_by_sku.get(str(sku).strip().casefold())
    if not master:
        unmatched[str(sku)] += 1
        continue
    matched += 1
    parts = [part.strip() for part in str(variation_name or "").split(",")]
    if parts and parts[0]:
        color_map[parts[0]][str(master["color"])] += 1
    if len(parts) > 1 and parts[-1]:
        size_map[parts[-1]][str(master["size"])] += 1
    key = str(product_name)
    product_map[key][f'{master["category"]}|{master["gender"]}|{master["sku"].split("-")[0]}'] += 1

print(json.dumps({
    "matched_existing_to_master": matched,
    "unmatched_existing_skus": unmatched,
    "color_label_to_master_color": {key: value for key, value in sorted(color_map.items())},
    "size_label_to_master_size": {key: value for key, value in sorted(size_map.items())},
    "product_to_master_type": {key: value for key, value in product_map.items()},
}, ensure_ascii=False, indent=2, default=lambda obj: dict(obj)))
