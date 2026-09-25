import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


root = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Shopee to tiktok")
source_path = next(root.glob("mass_update_sales*.xlsx"))
tiktok_paths = sorted(root.rglob("Tiktoksellercenter*.xlsx"))
output_dir = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok\inspection")
output_dir.mkdir(parents=True, exist_ok=True)


def sanitized_bytes(file_path):
    input_buffer = io.BytesIO(file_path.read_bytes())
    output_buffer = io.BytesIO()
    replacements = {
        b'activePane="bottom_left"': b'activePane="bottomLeft"',
        b'activePane="top_left"': b'activePane="topLeft"',
        b'activePane="bottom_right"': b'activePane="bottomRight"',
        b'activePane="top_right"': b'activePane="topRight"',
    }
    changed = False
    with zipfile.ZipFile(input_buffer, "r") as source_zip, zipfile.ZipFile(output_buffer, "w") as destination_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                for old, new in replacements.items():
                    if old in data:
                        data = data.replace(old, new)
                        changed = True
            destination_zip.writestr(item, data)
    return output_buffer.getvalue() if changed else file_path.read_bytes()


def norm_text(value):
    value = str(value or "").casefold()
    value = value.replace("tuxedo", "tuksedo").replace("tuskedo", "tuksedo")
    value = re.sub(r"[^0-9a-z]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


color_aliases = {
    "abu abu": "abu", "abu": "abu", "baby pink": "baby pink", "biru": "biru bca",
    "biru bca": "biru bca", "biru langit": "biru langit", "sky blue": "biru langit",
    "iced blue": "biru langit", "burgundy": "marun tua", "coklat": "coklat mahogani",
    "coklat mahogani": "coklat mahogani", "cream": "krem", "krem": "krem",
    "hitam": "hitam", "mahogani": "coklat mahogani", "maroon": "marun muda",
    "marun": "marun muda", "marun muda": "marun muda", "marun tua": "marun tua",
    "merah cabe": "merah cabe", "merah cabai": "merah cabe", "merah maroon": "marun muda",
    "navy": "navy", "navy blue": "navy", "pink fanta": "pink fanta", "putih": "putih",
    "sage": "sage", "sage hijau toska": "sage", "terracota": "terracota",
}


def norm_variation(value):
    raw = str(value or "").strip()
    if not raw:
        return ""
    parts = [part.strip() for part in raw.split(",")]
    normalized = []
    for index, part in enumerate(parts):
        token = norm_text(part)
        if index == 0:
            token = color_aliases.get(token, token)
        token = re.sub(r"\b(thn|t)\b", "tahun", token)
        normalized.append(token)
    return "|".join(normalized)


stopwords = {
    "best", "seller", "pakaian", "outfit", "style", "korean", "korea", "model", "premium",
    "bahan", "kualitas", "untuk", "dengan", "dan", "produk", "kantor", "meeting", "wisuda",
    "sidang", "skripsi", "nikah", "pengantin", "acara", "dewasa", "muda", "kekinian", "casual",
    "maupun", "slimfit", "slim", "fit", "original", "by", "fashion", "tren", "mode",
}


def title_tokens(value):
    return [token for token in norm_text(value).split() if token not in stopwords]


def title_scores(left, right):
    left_norm = " ".join(title_tokens(left))
    right_norm = " ".join(title_tokens(right))
    left_set, right_set = set(left_norm.split()), set(right_norm.split())
    jaccard = len(left_set & right_set) / max(1, len(left_set | right_set))
    sequence = SequenceMatcher(None, left_norm, right_norm).ratio()
    return jaccard, sequence


source_book = load_workbook(io.BytesIO(sanitized_bytes(source_path)), read_only=False, data_only=True)
source_sheet = source_book[source_book.sheetnames[0]]
source_groups = defaultdict(lambda: {"product_name": None, "rows": [], "variation_to_skus": defaultdict(set)})
source_sku_rows = 0
for row in range(7, source_sheet.max_row + 1):
    product_id = str(source_sheet.cell(row, 1).value or "").strip()
    product_name = source_sheet.cell(row, 2).value
    variation = source_sheet.cell(row, 4).value
    parent_sku = source_sheet.cell(row, 5).value
    variation_sku = source_sheet.cell(row, 6).value
    sku = variation_sku if variation_sku is not None and str(variation_sku).strip() else parent_sku
    if not product_id or sku is None or str(sku).strip() == "":
        continue
    sku = str(sku).strip()
    key = product_id
    group = source_groups[key]
    group["product_name"] = product_name
    group["rows"].append({"row": row, "variation": variation, "sku": sku})
    group["variation_to_skus"][norm_variation(variation)].add(sku)
    source_sku_rows += 1
source_book.close()

report = {
    "source_file": source_path.name,
    "source_sku_rows": source_sku_rows,
    "source_products_with_sku": len(source_groups),
    "source_groups": [
        {
            "source_product_id": product_id,
            "source_product_name": group["product_name"],
            "row_count": len(group["rows"]),
            "variation_count": len(group["variation_to_skus"]),
            "sku_prefixes": dict(Counter(row["sku"].split("-")[0].upper() for row in group["rows"])),
            "sample_rows": group["rows"][:8],
        }
        for product_id, group in source_groups.items()
    ],
    "tiktok_files": [],
}

for tiktok_path in tiktok_paths:
    workbook = load_workbook(io.BytesIO(sanitized_bytes(tiktok_path)), read_only=False, data_only=True)
    sheet = workbook["Template"]
    groups = defaultdict(lambda: {"product_name": None, "rows": [], "variations": set()})
    existing_seller_sku = 0
    for row in range(6, sheet.max_row + 1):
        product_id = str(sheet.cell(row, 1).value or "").strip()
        product_name = sheet.cell(row, 3).value
        variation = sheet.cell(row, 6).value
        seller_sku = sheet.cell(row, 11).value
        if not product_id:
            continue
        group = groups[product_id]
        group["product_name"] = product_name
        group["rows"].append({"row": row, "variation": variation, "seller_sku": seller_sku})
        group["variations"].add(norm_variation(variation))
        if seller_sku is not None and str(seller_sku).strip() != "":
            existing_seller_sku += 1

    group_reports = []
    for product_id, group in groups.items():
        candidates = []
        for source_product_id, source_group in source_groups.items():
            source_variations = set(source_group["variation_to_skus"])
            overlap = group["variations"] & source_variations
            jaccard, sequence = title_scores(group["product_name"], source_group["product_name"])
            coverage = len(overlap) / max(1, len(group["variations"]))
            score = coverage * 0.55 + jaccard * 0.30 + sequence * 0.15
            candidates.append({
                "source_product_id": source_product_id,
                "source_product_name": source_group["product_name"],
                "overlap": len(overlap),
                "tiktok_variants": len(group["variations"]),
                "source_variants": len(source_variations),
                "coverage": round(coverage, 4),
                "title_jaccard": round(jaccard, 4),
                "title_sequence": round(sequence, 4),
                "score": round(score, 4),
            })
        candidates.sort(key=lambda item: (item["score"], item["coverage"], item["title_jaccard"]), reverse=True)
        group_reports.append({
            "tiktok_product_id": product_id,
            "tiktok_product_name": group["product_name"],
            "row_count": len(group["rows"]),
            "existing_seller_sku": sum(1 for row in group["rows"] if row["seller_sku"] not in (None, "")),
            "variations": [row["variation"] for row in group["rows"]],
            "top_candidates": candidates[:5],
        })
    workbook.close()
    report["tiktok_files"].append({
        "file": tiktok_path.name,
        "data_rows": sum(len(group["rows"]) for group in groups.values()),
        "product_groups": len(groups),
        "existing_seller_sku": existing_seller_sku,
        "groups": group_reports,
    })

(output_dir / "matching_candidates.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({
    "source_sku_rows": source_sku_rows,
    "source_products_with_sku": len(source_groups),
    "tiktok_files": [
        {key: value for key, value in item.items() if key != "groups"}
        for item in report["tiktok_files"]
    ],
}, ensure_ascii=False, indent=2))
