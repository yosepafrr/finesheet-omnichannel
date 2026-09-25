import csv
import json
from pathlib import Path

from openpyxl import load_workbook


input_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")
updates_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\inspection\proposed_sku_updates.csv")
output_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\mass_update_sales_info_615078427_20260925103824_SKU_updated.xlsx")

with updates_path.open("r", encoding="utf-8-sig", newline="") as handle:
    updates = list(csv.DictReader(handle))

workbook = load_workbook(input_path, read_only=False, data_only=False)
sheet = workbook[workbook.sheetnames[0]]

before_existing = {}
before_blank_rows = []
for row_number in range(7, sheet.max_row + 1):
    value = sheet.cell(row_number, 6).value
    if value is None or str(value).strip() == "":
        before_blank_rows.append(row_number)
    else:
        before_existing[row_number] = value

applied = []
for update in updates:
    row_number = int(update["row"])
    cell = sheet.cell(row_number, 6)
    if cell.value is not None and str(cell.value).strip() != "":
        raise RuntimeError(f"Refusing to overwrite populated SKU at F{row_number}: {cell.value!r}")
    cell.value = update["sku"]
    applied.append(row_number)

output_path.parent.mkdir(parents=True, exist_ok=True)
workbook.save(output_path)
workbook.close()

result = {
    "output": str(output_path),
    "applied_count": len(applied),
    "blank_before": len(before_blank_rows),
    "expected_blank_after": len(before_blank_rows) - len(applied),
    "preserved_existing_count": len(before_existing),
    "first_updated_rows": applied[:10],
    "last_updated_rows": applied[-10:],
}
print(json.dumps(result, ensure_ascii=False, indent=2))
