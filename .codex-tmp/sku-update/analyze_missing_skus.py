import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


template_path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")
output_dir = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\inspection")
output_dir.mkdir(parents=True, exist_ok=True)

workbook = load_workbook(template_path, read_only=False, data_only=False)
sheet = workbook[workbook.sheetnames[0]]

rows = []
for row_number in range(7, sheet.max_row + 1):
    values = [sheet.cell(row_number, column).value for column in range(1, 16)]
    product_id, product_name, variation_id, variation_name, parent_sku, sku = values[:6]
    if not product_id:
        continue
    rows.append(
        {
            "row": row_number,
            "product_id": str(product_id),
            "product_name": product_name,
            "variation_id": str(variation_id) if variation_id is not None else None,
            "variation_name": variation_name,
            "parent_sku": parent_sku,
            "sku": sku,
        }
    )

missing = [row for row in rows if row["sku"] is None or str(row["sku"]).strip() == ""]
by_product = defaultdict(lambda: {"missing": 0, "filled": 0, "samples_missing": [], "samples_filled": []})
for row in rows:
    bucket = by_product[(row["product_id"], row["product_name"])]
    if row["sku"] is None or str(row["sku"]).strip() == "":
        bucket["missing"] += 1
        if len(bucket["samples_missing"]) < 12:
            bucket["samples_missing"].append({"row": row["row"], "variation": row["variation_name"]})
    else:
        bucket["filled"] += 1
        if len(bucket["samples_filled"]) < 12:
            bucket["samples_filled"].append({"row": row["row"], "variation": row["variation_name"], "sku": row["sku"]})

summary = {
    "total_data_rows": len(rows),
    "filled_sku_rows": len(rows) - len(missing),
    "missing_sku_rows": len(missing),
    "products_with_missing": [
        {
            "product_id": key[0],
            "product_name": key[1],
            **value,
        }
        for key, value in by_product.items()
        if value["missing"]
    ],
}

(output_dir / "missing_sku_summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
)
with (output_dir / "missing_sku_rows.csv").open("w", newline="", encoding="utf-8-sig") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(missing[0].keys()) if missing else ["row"])
    writer.writeheader()
    writer.writerows(missing)

print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
