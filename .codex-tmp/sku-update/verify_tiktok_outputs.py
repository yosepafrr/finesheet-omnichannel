import csv
import hashlib
import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIR = Path(r"C:\Bisnis\Vilion Apparel\Products SKUs\Shopee to tiktok\Tiktoksellercenter_batchedit_20260925_all_information_template")
UPDATES_PATH = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok\inspection\proposed_tiktok_sku_updates.csv")
OUTPUT_DIR = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\shopee-to-tiktok")
REPORT_PATH = OUTPUT_DIR / "verification_report.json"
SHEET_XML = "xl/worksheets/sheet1.xml"


def cell_pattern(coordinate: str) -> re.Pattern[str]:
    return re.compile(
        rf'<c\b(?=[^>]*\br="{re.escape(coordinate)}")[^>]*(?:/>|>.*?</c>)',
        flags=re.DOTALL,
    )


updates_by_file: dict[str, list[dict[str, str]]] = defaultdict(list)
with UPDATES_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
    for row in csv.DictReader(handle):
        updates_by_file[row["file"]].append(row)

results = []
total_updates = 0
total_remaining_blank = 0
for source_path in sorted(SOURCE_DIR.glob("*.xlsx")):
    rows = sorted(updates_by_file.get(source_path.name, []), key=lambda item: int(item["row"]))
    output_path = OUTPUT_DIR / f"{source_path.stem}_SKU_updated.xlsx"
    if not output_path.exists():
        raise RuntimeError(f"Missing output: {output_path}")

    before_book = load_workbook(source_path, read_only=True, data_only=False)
    after_book = load_workbook(output_path, read_only=True, data_only=False)
    before_sheet = before_book["Template"]
    after_sheet = after_book["Template"]
    assert before_book.sheetnames == after_book.sheetnames
    # TikTok's declared worksheet dimension ends at row 5 even though data rows
    # follow it, so derive the real final row from worksheet XML.
    with zipfile.ZipFile(source_path, "r") as dimension_zip:
        source_sheet_xml = dimension_zip.read(SHEET_XML).decode("utf-8")
    data_max_row = max(int(value) for value in re.findall(r'<row\b[^>]*\br="(\d+)"', source_sheet_xml))

    approved_rows = {int(item["row"]): item for item in rows}
    changed_rows = set()
    remaining_blank = 0
    before_rows = before_sheet.iter_rows(min_row=6, max_row=data_max_row, min_col=1, max_col=11, values_only=True)
    after_rows = after_sheet.iter_rows(min_row=6, max_row=data_max_row, min_col=1, max_col=11, values_only=True)
    for row_number, (before_values, after_values) in enumerate(zip(before_rows, after_rows), start=6):
        before_value = before_values[10]
        after_value = after_values[10]
        if row_number in approved_rows:
            update = approved_rows[row_number]
            assert before_value is None or str(before_value).strip() == ""
            assert after_value == update["sku"], (source_path.name, row_number, after_value, update["sku"])
            assert str(before_values[0]) == update["product_id"]
            assert str(before_values[5]) == update["variation"]
            changed_rows.add(row_number)
        else:
            assert after_value == before_value, (source_path.name, row_number, before_value, after_value)
        if after_value is None or str(after_value).strip() == "":
            remaining_blank += 1

    assert changed_rows == set(approved_rows)

    with zipfile.ZipFile(source_path, "r") as before_zip, zipfile.ZipFile(output_path, "r") as after_zip:
        assert before_zip.namelist() == after_zip.namelist()
        if not rows:
            assert source_path.read_bytes() == output_path.read_bytes(), f"Zero-update file changed: {source_path.name}"
            unchanged_entries = len(before_zip.namelist())
        else:
            for name in before_zip.namelist():
                if name != SHEET_XML:
                    assert before_zip.read(name) == after_zip.read(name), f"Unexpected package change: {source_path.name}: {name}"
            source_xml = before_zip.read(SHEET_XML).decode("utf-8")
            restored_xml = after_zip.read(SHEET_XML).decode("utf-8")
            for update in rows:
                coordinate = f"K{update['row']}"
                pattern = cell_pattern(coordinate)
                source_match = pattern.search(source_xml)
                output_match = pattern.search(restored_xml)
                assert source_match and output_match, (source_path.name, coordinate)
                assert f"<t>{update['sku']}</t>" in output_match.group(0), (source_path.name, coordinate)
                restored_xml = restored_xml[: output_match.start()] + source_match.group(0) + restored_xml[output_match.end() :]
            assert restored_xml == source_xml, f"Worksheet changed outside approved cells: {source_path.name}"
            unchanged_entries = len(before_zip.namelist()) - 1

    result = {
        "file": output_path.name,
        "verified_updates": len(rows),
        "remaining_blank_skus": remaining_blank,
        "unchanged_package_entries": unchanged_entries,
        "sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "size_bytes": output_path.stat().st_size,
    }
    results.append(result)
    total_updates += len(rows)
    total_remaining_blank += remaining_blank

report = {
    "verified_total_updates": total_updates,
    "verified_total_remaining_blank_skus": total_remaining_blank,
    "unexpected_value_diffs": 0,
    "unexpected_package_changes": 0,
    "files": results,
}
REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
