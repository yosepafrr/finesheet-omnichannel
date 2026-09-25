import json
from pathlib import Path

from openpyxl import load_workbook


path = Path(r"C:\Bisnis\finesheet-omnichannel\outputs\sku-update\working\mass_update_sales_sanitized.xlsx")
workbook = load_workbook(path, read_only=False, data_only=False)
for sheet in workbook.worksheets:
    print(json.dumps({
        "sheet": sheet.title,
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "calculated_dimension": sheet.calculate_dimension(),
    }, ensure_ascii=False))
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), max_col=min(sheet.max_column, 50), values_only=True):
        print(json.dumps(list(row), ensure_ascii=False, default=str))
